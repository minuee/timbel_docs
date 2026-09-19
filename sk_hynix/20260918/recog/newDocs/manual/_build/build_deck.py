#!/usr/bin/env python3
"""16:9 static step boards + the PowerPoint deck."""
import os, sys
sys.path.insert(0, "/home/claude/manual/work")
os.chdir("/home/claude/manual")
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from steps import UNITS, ORDER, SECTIONS, STEP_CROP
try:
    from steps import BOARD_COLS
except ImportError:
    BOARD_COLS = {}
from frames_lib import processed
from build_stills import wrap

BOLD = "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"
REG  = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
fnt  = lambda p, s: ImageFont.truetype(p, s)

INK, GRAY, ACCENT, LINE = (26, 26, 26), (110, 116, 124), (232, 89, 12), (223, 227, 232)
SW, SH = 1920, 1080


# ---------------------------------------------------------------- boards ---
def board(unit):
    meta = UNITS[unit]
    frames = processed(unit)
    crop = STEP_CROP.get(unit)
    shots = []
    for st in meta["steps"]:
        c = st[2] if len(st) > 2 else crop
        img = frames[min(st[0], len(frames) - 1)]
        shots.append(img.crop(c) if c else img)

    n = len(shots)
    portrait = shots[0].height > shots[0].width * 1.1
    cols = BOARD_COLS.get(unit) or (n if portrait else (2 if n <= 4 else 3))
    rows = (n + cols - 1) // cols

    canvas = Image.new("RGB", (SW, SH), "white")
    d = ImageDraw.Draw(canvas)
    d.rectangle([0, 0, SW, 10], fill=ACCENT)

    ftitle, fintro = fnt(BOLD, 52), fnt(REG, 26)
    fcap, fnum = fnt(REG, 25), fnt(BOLD, 24)

    d.text((70, 52), meta["title"], fill=INK, font=ftitle)
    d.text((70, 124), meta["intro"], fill=GRAY, font=fintro)

    top, pad = 190, 40
    cw = (SW - 2 * 70 - (cols - 1) * pad) // cols
    ch = (SH - top - 70 - (rows - 1) * pad) // rows

    for i, shot in enumerate(shots):
        cx = 70 + (i % cols) * (cw + pad)
        cy = top + (i // cols) * (ch + pad)
        lines = wrap(d, meta["steps"][i][1], fcap, cw - 56)[:3]
        head = len(lines) * 32 + 12
        r = 17
        d.ellipse([cx, cy + 2, cx + 2 * r, cy + 2 + 2 * r], fill=ACCENT)
        tw = d.textlength(str(i + 1), font=fnum)
        d.text((cx + r - tw / 2, cy + 3), str(i + 1), fill="white", font=fnum)
        for j, ln in enumerate(lines):
            d.text((cx + 2 * r + 12, cy + 2 + j * 32), ln, fill=INK, font=fcap)

        avail_h, avail_w = ch - head, cw
        s = min(avail_w / shot.width, avail_h / shot.height)
        t = shot.resize((max(1, int(shot.width * s)), max(1, int(shot.height * s))), Image.LANCZOS)
        ox = cx + (cw - t.width) // 2
        oy = cy + head
        d.rectangle([ox - 1, oy - 1, ox + t.width, oy + t.height], outline=LINE)
        canvas.paste(t, (ox, oy))

    os.makedirs("slide", exist_ok=True)
    canvas.save(f"slide/{unit}_board.png")
    return f"slide/{unit}_board.png"


# ------------------------------------------------------------------ deck ---
SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)


def txt(slide, x, y, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, space=0):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_after = Pt(space)
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = "맑은 고딕"
            r.font.color.rgb = RGBColor(*color)
    return tb


def bar(slide, color=ACCENT, h=Inches(0.09)):
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, h)
    s.fill.solid(); s.fill.fore_color.rgb = RGBColor(*color)
    s.line.fill.background(); s.shadow.inherit = False


def fit_picture(slide, path, left, top, max_w, max_h):
    with Image.open(path) as im:
        w, h = im.size
    s = min(max_w / w, max_h / h)
    pw, ph = int(w * s), int(h * s)
    return slide.shapes.add_picture(path, left + int((max_w - pw) / 2),
                                    top + int((max_h - ph) / 2), pw, ph)


def build_deck(out, gif_dir="gif"):
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    blank = prs.slide_layouts[6]

    # title ---------------------------------------------------------------
    s = prs.slides.add_slide(blank); bar(s, ACCENT, Inches(0.14))
    txt(s, Inches(1.1), Inches(2.4), Inches(11), Inches(1.2), "AI 회의록 사용 안내", 54, True)
    txt(s, Inches(1.1), Inches(3.5), Inches(11), Inches(0.8),
        "폴더 관리 · 녹음과 AI 요약", 24, False, GRAY)
    txt(s, Inches(1.1), Inches(6.3), Inches(11), Inches(0.5),
        "슬라이드쇼(F5)로 보시면 화면 동작이 재생됩니다.", 14, False, GRAY)

    # how to read ---------------------------------------------------------
    s = prs.slides.add_slide(blank); bar(s)
    txt(s, Inches(0.9), Inches(0.6), Inches(11.5), Inches(0.8), "이 문서를 보는 방법", 36, True)
    body = ("• 기능마다 두 장씩 있습니다 — 움직이는 화면 1장, 정지 화면 1장\n"
            "• 움직이는 화면은 슬라이드쇼(F5)에서만 재생됩니다. 편집 화면에서는 첫 장면만 보입니다\n"
            "• PDF로 저장하거나 인쇄하면 움직임은 남지 않습니다. 이때는 정지 화면 쪽을 보세요")
    txt(s, Inches(0.9), Inches(1.9), Inches(11.5), Inches(3), body, 20, False, INK, space=14)

    for sect, units in SECTIONS:
        # section divider -------------------------------------------------
        s = prs.slides.add_slide(blank); bar(s, ACCENT, Inches(0.14))
        txt(s, Inches(1.1), Inches(3.0), Inches(11), Inches(1), sect, 44, True)
        txt(s, Inches(1.1), Inches(3.9), Inches(11), Inches(0.6),
            " · ".join(UNITS[u]["title"] for u in units), 20, False, GRAY)

        for u in units:
            meta = UNITS[u]
            # animated slide ----------------------------------------------
            s = prs.slides.add_slide(blank); bar(s)
            txt(s, Inches(0.55), Inches(0.42), Inches(4.4), Inches(0.8), meta["title"], 32, True)
            lines = "\n".join(f"{i}.  {st[1]}" for i, st in enumerate(meta["steps"], 1))
            txt(s, Inches(0.55), Inches(1.5), Inches(4.1), Inches(5.4), lines, 15, False, INK, space=12)
            txt(s, Inches(0.55), Inches(6.85), Inches(4.1), Inches(0.4),
                "▶ 슬라이드쇼에서 재생됩니다", 11, False, GRAY)
            fit_picture(s, f"{gif_dir}/{u}.gif", Inches(4.95), Inches(0.62),
                        Inches(7.85), Inches(6.3))

            # static slide -------------------------------------------------
            s = prs.slides.add_slide(blank)
            s.shapes.add_picture(f"slide/{u}_board.png", 0, 0, SLIDE_W, SLIDE_H)

    prs.save(out)
    print("saved", out, round(os.path.getsize(out) / 1e6, 2), "MB")


if __name__ == "__main__":
    for u in ORDER:
        print("board:", board(u))
    os.makedirs("ppt", exist_ok=True)
    build_deck("ppt/AI회의록_사용안내.pptx", "gif")
    build_deck("ppt/AI회의록_사용안내_경량.pptx", "gif-light")
