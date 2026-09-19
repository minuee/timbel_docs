#!/usr/bin/env python3
"""The full guide deck: cover, contents, 6 parts, 23 units."""
import os, sys
sys.path.insert(0, "/home/claude/manual/work")
os.chdir("/home/claude/manual")
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

from steps import UNITS
from guide import SECTIONS, PLACEHOLDERS, ALL_UNITS, REAL_UNITS, NUMBER, title_of
from build_deck import board, txt, bar, fit_picture, INK, GRAY, ACCENT, LINE, SLIDE_W, SLIDE_H

DASH = (176, 182, 190)


def placeholder_slide(prs, blank, unit):
    ph = PLACEHOLDERS[unit]
    s = prs.slides.add_slide(blank); bar(s)
    txt(s, Inches(0.55), Inches(0.42), Inches(8), Inches(0.8),
        f"{NUMBER[unit]:02d}.  {ph['title']}", 32, True)

    badge = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                               Inches(0.58), Inches(1.28), Inches(1.15), Inches(0.38))
    badge.fill.solid(); badge.fill.fore_color.rgb = RGBColor(*ACCENT)
    badge.line.fill.background(); badge.shadow.inherit = False
    tf = badge.text_frame; tf.text = "준비 중"
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    p.runs[0].font.size = Pt(14); p.runs[0].font.bold = True
    p.runs[0].font.name = "맑은 고딕"; p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

    txt(s, Inches(0.55), Inches(1.95), Inches(5.2), Inches(0.5),
        "이 단원에는 다음 내용이 들어갑니다", 17, True, GRAY)
    txt(s, Inches(0.55), Inches(2.5), Inches(5.2), Inches(2.5),
        "\n".join(f"·  {t}" for t in ph["todo"]), 17, False, INK, space=14)
    txt(s, Inches(0.55), Inches(6.5), Inches(5.2), Inches(0.7),
        "화면 녹화가 준비되면 오른쪽 자리에\n넣어 주세요.", 13, False, GRAY)

    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                             Inches(6.2), Inches(1.3), Inches(6.5), Inches(5.6))
    box.fill.solid(); box.fill.fore_color.rgb = RGBColor(250, 250, 251)
    box.line.color.rgb = RGBColor(*DASH); box.line.width = Pt(1.5)
    box.line.dash_style = 4  # dashed
    box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.text = "화면 캡처 / GIF 자리"
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    p.runs[0].font.size = Pt(18); p.runs[0].font.bold = True
    p.runs[0].font.name = "맑은 고딕"; p.runs[0].font.color.rgb = RGBColor(*DASH)


def unit_slides(prs, blank, unit, gif_dir):
    meta = UNITS[unit]
    s = prs.slides.add_slide(blank); bar(s)
    tsize = 28 if len(meta["title"]) <= 9 else 23
    txt(s, Inches(0.55), Inches(0.42), Inches(4.3), Inches(0.9),
        f"{NUMBER[unit]:02d}.  {meta['title']}", tsize, True)
    lines = "\n".join(f"{i}.  {st[1]}" for i, st in enumerate(meta["steps"], 1))
    n = len(meta["steps"])
    fs, sp = (14, 12) if n <= 5 else ((12, 8) if n <= 7 else (11, 4))
    txt(s, Inches(0.55), Inches(1.55), Inches(4.1), Inches(5.3), lines, fs, False, INK, space=sp)
    txt(s, Inches(0.55), Inches(6.85), Inches(4.1), Inches(0.4),
        "▶ 슬라이드쇼에서 재생됩니다", 11, False, GRAY)
    fit_picture(s, f"{gif_dir}/{unit}.gif", Inches(4.95), Inches(0.62),
                Inches(7.85), Inches(6.3))

    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(f"slide/{unit}_board.jpg", 0, 0, SLIDE_W, SLIDE_H)


def contents_slide(prs, blank):
    s = prs.slides.add_slide(blank); bar(s)
    txt(s, Inches(0.7), Inches(0.5), Inches(11.5), Inches(0.8), "목차", 34, True)
    cols = [[], []]
    flat = []
    for sect, units in SECTIONS:
        flat.append(("S", sect))
        for u in units:
            flat.append(("U", f"{NUMBER[u]:02d}  {title_of(u, UNITS)}"
                              + ("   〈준비 중〉" if u.startswith("P") else "")))
    half = 16
    cols[0], cols[1] = flat[:half], flat[half:]
    for ci, col in enumerate(cols):
        x = Inches(0.8 + ci * 6.2)
        y = Inches(1.35)
        for kind, text in col:
            if kind == "S":
                txt(s, x, y, Inches(5.6), Inches(0.4), text, 15, True, ACCENT)
                y += Inches(0.42)
            else:
                txt(s, x + Inches(0.15), y, Inches(5.5), Inches(0.35), text, 13, False, INK)
                y += Inches(0.32)


def build(out, gif_dir="gif"):
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank); bar(s, ACCENT, Inches(0.14))
    txt(s, Inches(1.1), Inches(2.3), Inches(11), Inches(1.3), "AI 회의록 전체 가이드", 54, True)
    txt(s, Inches(1.1), Inches(3.45), Inches(11), Inches(0.8),
        "시작하기 · 회의록 만들기 · 확인하기 · 정리하기 · 찾기 · 관리와 설정", 22, False, GRAY)
    txt(s, Inches(1.1), Inches(6.3), Inches(11), Inches(0.5),
        "슬라이드쇼(F5)로 보시면 화면 동작이 재생됩니다.", 14, False, GRAY)

    contents_slide(prs, blank)

    s = prs.slides.add_slide(blank); bar(s)
    txt(s, Inches(0.9), Inches(0.6), Inches(11.5), Inches(0.8), "이 문서를 보는 방법", 36, True)
    body = ("• 기능마다 두 장씩 있습니다 — 움직이는 화면 1장, 정지 화면 1장\n"
            "• 움직이는 화면은 슬라이드쇼(F5)에서만 재생됩니다. 편집 화면에서는 첫 장면만 보입니다\n"
            "• PDF로 저장하거나 인쇄하면 움직임은 남지 않습니다. 이때는 정지 화면 쪽을 보세요\n"
            "• 〈준비 중〉으로 표시된 단원은 화면 녹화가 준비되는 대로 채워집니다")
    txt(s, Inches(0.9), Inches(1.9), Inches(11.5), Inches(3), body, 20, False, INK, space=14)

    for sect, units in SECTIONS:
        s = prs.slides.add_slide(blank); bar(s, ACCENT, Inches(0.14))
        txt(s, Inches(1.1), Inches(3.0), Inches(11), Inches(1), sect, 44, True)
        txt(s, Inches(1.1), Inches(3.95), Inches(11), Inches(0.9),
            " · ".join(title_of(u, UNITS) for u in units), 18, False, GRAY)
        for u in units:
            if u.startswith("P"):
                placeholder_slide(prs, blank, u)
            else:
                unit_slides(prs, blank, u, gif_dir)

    prs.save(out)
    print("saved", out, len(prs.slides.__iter__.__self__._sldIdLst), "slides",
          round(os.path.getsize(out) / 1e6, 2), "MB")


if __name__ == "__main__":
    for u in REAL_UNITS:
        board(u)
    os.makedirs("ppt", exist_ok=True)
    build("ppt/AI회의록_전체가이드.pptx", "gif-mid")
    build("ppt/AI회의록_전체가이드_경량.pptx", "gif-light")
